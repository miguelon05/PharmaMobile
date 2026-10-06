from pathlib import Path
from html.parser import HTMLParser
import subprocess
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
src=Path('docs/generar_manual.py').read_text(encoding='utf-8-sig'); ns={}; exec(src[:src.index('doc=Doc(str(OUT)')],ns)
qa=Path('tmp/word-qa'); qa.mkdir(parents=True,exist_ok=True)
c=canvas.Canvas(str(qa/'diagram.pdf'),pagesize=(499,295)); ns['Diagram']().drawOn(c,0,0); c.save()
subprocess.run(['pdftoppm','-r','180','-singlefile','-png',str(qa/'diagram.pdf'),str(qa/'diagram')],check=True)
d=Document(); sec=d.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7); sec.top_margin=Cm(1.7); sec.bottom_margin=Cm(1.9); sec.left_margin=sec.right_margin=Cm(1.7)
for st in d.styles:
 if st.type==1 or st.type==2:
  st.font.name='Arial'; st.font.color.rgb=RGBColor(0,0,0)
for name,size in [('Normal',10),('Title',24),('Heading 1',17),('Heading 2',11),('Caption',8.5)]:
 s=d.styles[name]; s.font.size=Pt(size); s.paragraph_format.space_after=Pt(8); s.paragraph_format.line_spacing=1.1
for name in ['Title','Heading 1','Heading 2']: d.styles[name].font.bold=True
class Runs(HTMLParser):
 def __init__(self,p): super().__init__(); self.p=p; self.bold=False; self.italic=False
 def handle_starttag(self,t,a):
  if t=='b': self.bold=True
  if t=='i': self.italic=True
  if t=='br': self.p.add_run().add_break()
 def handle_endtag(self,t):
  if t=='b': self.bold=False
  if t=='i': self.italic=False
 def handle_data(self,s):
  r=self.p.add_run(s); r.bold=self.bold or self.p.style.name in ['Title','Heading 1','Heading 2']; r.italic=self.italic; r.font.color.rgb=RGBColor(0,0,0)
def fill(p,t): Runs(p).feed(t)
def table(flow):
 vals=flow._cellvalues; t=d.add_table(rows=len(vals),cols=len(vals[0])); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
 widths=flow._colWidths
 for i,w in enumerate(widths): t.columns[i].width=Pt(w)
 pr=t._tbl.tblPr; borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  x=OxmlElement('w:'+side); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'5'); x.set(qn('w:color'),'000000'); borders.append(x)
 pr.append(borders)
 for i,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr(); no=OxmlElement('w:cantSplit'); trpr.append(no)
  if i==0 and len(vals)>1: trpr.append(OxmlElement('w:tblHeader'))
  for j,cell in enumerate(row.cells):
   cell.width=Pt(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side); e.set(qn('w:w'),'85'); e.set(qn('w:type'),'dxa'); mar.append(e)
   cell._tc.get_or_add_tcPr().append(mar)
   p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.05
   value=vals[i][j]; fill(p,value.text if hasattr(value,'text') else str(value))
   for r in p.runs: r.font.size=Pt(8); r.bold=(i==0 and len(vals)>1) or r.bold
 if len(vals)==1:
  t.rows[0].height=Pt(flow._argH[0]); t.rows[0].height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
  t.cell(0,0).vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
  p=t.cell(0,0).add_paragraph('Pega aquí la imagen y elimina estas indicaciones.'); p.paragraph_format.space_before=Pt(12)
  for r in p.runs:r.font.size=Pt(9)
 d.add_paragraph().paragraph_format.space_after=Pt(0)
for f in ns['story']:
 name=type(f).__name__
 if name=='Paragraph':
  sn=f.style.name; text=f.text
  sty={'TitleBlack':'Title','SectionBlack':'Heading 1','SubBlack':'Heading 2','SmallBlack':'Caption'}.get(sn,'Normal')
  if text=='Responsabilidades por capa y pruebas de flujo':sty='Subtitle'
  
  if text=='Clases de pruebas del módulo shared': d.add_page_break()
  p=d.add_paragraph(style=sty); fill(p,text)
 elif name=='PageBreak': d.add_page_break()
 elif name=='Spacer':
  if f.height>20:
   p=d.add_paragraph(); p.paragraph_format.space_after=Pt(f.height-12)
 elif name=='Table': table(f)
 elif name=='Diagram': d.add_picture(str(qa/'diagram.png'),width=Pt(499))
 elif name=='TableOfContents':
  for s in ['1. Mapa de capas','2. Tabla de responsabilidades','3. Decisiones justificadas','4. Pruebas de flujo','5. Bitácora de dificultades']: d.add_paragraph(s)
f=sec.footer.paragraphs[0]; f.alignment=2
r=f.add_run('PharmaMobil | '); r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)
d.core_properties.title='Manual de arquitectura de PharmaMobil'; d.core_properties.author='Acevedo'
for el in list(d.styles.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
for el in list(d.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
for el in d.styles.element.iter(qn('w:color')):
 el.set(qn('w:val'),'000000')
 for attr in ['themeColor','themeTint','themeShade']:
  el.attrib.pop(qn('w:'+attr),None)
out=Path('output/docx/Manual-Arquitectura-PharmaMobil_Acevedo.docx'); out.parent.mkdir(parents=True,exist_ok=True); d.save(out); print(out)
